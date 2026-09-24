# -*- coding: utf-8 -*-
"""B2 知识库管理模块 · 入库全链路单元测试
运行: cd backend && python tests/test_ingest.py
覆盖: 5 格式提取 / 清洗 / 分块 / 幂等 upsert / id与元数据 / 删除+缓存失效 / 清库
"""
import hashlib
import os
import sys
import tempfile
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

TMP = tempfile.mkdtemp(prefix="gk-ingest-test-")
os.environ["DB_PATH"] = os.path.join(TMP, "vec")
os.environ["DATA_DIR"] = os.path.join(TMP, "data")
os.environ["MODEL_CACHE"] = os.path.join(TMP, "data", "models")
os.environ["EMBED_PROVIDER"] = "selftest"

ok = True


def check(name, cond, extra=""):
    global ok
    print(("PASS " if cond else "FAIL ") + name + (("  " + extra) if extra and not cond else ""))
    if not cond:
        ok = False


from poc.config import Config  # noqa: E402
from poc.ingest import clean_text, chunk_text, extract_text, ingest_file  # noqa: E402
from poc import store  # noqa: E402

cfg = Config()
TID = "T900"
KB_DIR = Path(TMP) / "kb"
KB_DIR.mkdir(parents=True, exist_ok=True)


def make_samples():
    md = KB_DIR / "讲义.md"
    md.write_text("# 转折关系\n转折之后是重点，转折之前可略读。\n隐性转折词包括实际上、事实上、其实。", encoding="utf-8")
    txt = KB_DIR / "讲义.txt"
    txt.write_text("逻辑判断题型分类：翻译推理、真假推理、分析推理。\n做题顺序：先翻译、再推理。", encoding="utf-8")
    docx_path = KB_DIR / "讲义.docx"
    import docx
    d = docx.Document()
    d.add_paragraph("公考逻辑判断：假言命题推理规则，肯前必肯后，否后必否前。")
    d.add_paragraph("逆否命题与原命题等价。")
    d.save(str(docx_path))
    pptx_path = KB_DIR / "讲义.pptx"
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    box.text_frame.text = "公考言语理解：转折关系核心考点"
    prs.save(str(pptx_path))
    pdf_path = KB_DIR / "讲义.pdf"
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Public exam logical reasoning PDF sample text for retrieval.")
    doc.save(str(pdf_path))
    doc.close()
    return md, txt, docx_path, pptx_path, pdf_path


MD, TXT, DOCX, PPTX, PDF = make_samples()
SAMPLES = [(MD, "md"), (TXT, "txt"), (DOCX, "docx"), (PPTX, "pptx"), (PDF, "pdf")]

# ---------- 1) 5 种格式提取到可见文本 ----------
for p, fmt in SAMPLES:
    chunks = ingest_file(p, cfg)
    total = sum(len(c.content.strip()) for c in chunks)
    check(f"格式[{fmt}] 提取可见文本", len(chunks) > 0 and total > 0, f"chunks={len(chunks)} chars={total}")

# ---------- 2) 清洗 ----------
check("清洗: \\r\\n→\\n 去空行 strip", clean_text("  a\r\n\r\n b \r\n") == "a\nb")
check("清洗: 纯空白→空", clean_text("  \n \n") == "")

# ---------- 3) 分块：段落聚合 + 重叠 ----------
para_text = "\n".join("段落%d：" % i + "公考知识点内容填充。  " * 20 for i in range(1, 6))
parts = chunk_text(para_text, 600, 120)
check("分块: 多段聚合数>1", len(parts) > 1, f"chunks={len(parts)}")
check("分块: 每块<=600", all(len(x) <= 600 for x in parts), [len(x) for x in parts])
check("分块: 相邻块重叠120", len(parts) >= 2 and parts[0][-120:] in parts[1])

# ---------- 3b) 单段超长硬切 ----------
long_para = "硬切测试。" * 200  # 1000 字单段
hard = chunk_text(long_para, 600, 120)
check("硬切: 单段拆成多块且每块<=600", len(hard) > 1 and all(len(x) <= 600 for x in hard), [len(x) for x in hard])
check("硬切: 拼接还原原文", "".join(hard) == long_para)

# ---------- 4) 入库：幂等 upsert / id / 元数据 ----------
store.clear_collection(TID)
big = KB_DIR / "长讲义.md"
big.write_text("公考言语理解转折关系核心知识点。\n" + "转折之后是重点，转折之前可略读。\n" * 80, encoding="utf-8")
chunks = ingest_file(big, cfg)
n1 = store.upsert_chunks(chunks, TID)
count1 = store.collection_count(TID)
check("入库: 首次写入块数一致", n1 == len(chunks) == count1, f"n1={n1} count={count1}")

n2 = store.upsert_chunks(chunks, TID)
count2 = store.collection_count(TID)
check("幂等: 重复入库块数不翻倍", n2 == len(chunks) and count2 == count1, f"count1={count1} count2={count2}")

coll = store._collection(TID)
data = coll.get(include=["metadatas"])
got_ids = sorted(data["ids"])
expect_ids = sorted(
    f"{c.doc_name}_{c.idx}_{hashlib.md5(c.content.encode('utf-8')).hexdigest()[:8]}" for c in chunks
)
check("id 格式: {doc_name}_{idx}_{md5[:8]}", got_ids == expect_ids,
      f"ids={len(got_ids)} expect={len(expect_ids)}")
meta_ok = all(
    m.get("teacher_id") == TID and m.get("doc_name") == big.name and isinstance(m.get("idx"), int)
    for m in data["metadatas"]
)
check("元数据: {teacher_id, doc_name, idx}", meta_ok, str(data["metadatas"][:1]))

docs = store.list_documents(TID)
check("list_documents: 按 doc_name 聚合", any(d["doc_name"] == big.name and d["chunks"] == len(chunks) for d in docs),
      str(docs))

# ---------- 5) 缓存失效：写入后紧接检索能命中新结果 ----------
store.upsert_chunks(ingest_file(MD, cfg), TID)
hits_new = store.retrieve("转折之后是重点", TID)
check("写入后紧接检索命中新知识", any(h.get("meta", {}).get("doc_name") == MD.name for h in hits_new),
      str([h.get("meta", {}).get("doc_name") for h in hits_new]))

# ---------- 6) 删除文档 + 缓存失效 ----------
store.retrieve("转折关系", TID)  # 预热缓存（TID 相关键已存在）
before = len([k for k in store.retrieve_cache._data if str(k).startswith(f"retr|{TID}|")])
deleted = store.delete_document(TID, MD.name)
after = len([k for k in store.retrieve_cache._data if str(k).startswith(f"retr|{TID}|")])
check("删除: 返回被删块数>0", deleted > 0, f"deleted={deleted}")
check("删除: 该师检索缓存已失效", before > 0 and after == 0, f"before={before} after={after}")
check("删除: list_documents 无该项", not any(d["doc_name"] == MD.name for d in store.list_documents(TID)))
hits_after = store.retrieve("转折关系", TID)
check("删除后紧接检索无旧结果", not any(h.get("meta", {}).get("doc_name") == MD.name for h in hits_after),
      str([h.get("meta", {}).get("doc_name") for h in hits_after]))

# ---------- 7) 异常：空文件 / 不支持格式 / 解析失败 ----------
empty_p = KB_DIR / "空文件.md"
empty_p.write_text("", encoding="utf-8")
try:
    ingest_file(empty_p, cfg)
    check("空文件: 抛 ValueError", False, "未抛异常")
except ValueError as e:
    check("空文件: 抛 ValueError(入库失败: ...)", "入库失败" in str(e), str(e))

bad_p = KB_DIR / "未知格式.xyz"
bad_p.write_text("随便什么内容", encoding="utf-8")
try:
    ingest_file(bad_p, cfg)
    check("不支持格式: 抛 ValueError", False, "未抛异常")
except ValueError as e:
    check("不支持格式: 抛 ValueError(入库失败: ...)", "入库失败" in str(e), str(e))

corrupt_p = KB_DIR / "坏文件.docx"
corrupt_p.write_bytes(b"not a real docx file")
try:
    ingest_file(corrupt_p, cfg)
    check("坏文件(解析失败): 抛 ValueError", False, "未抛异常")
except ValueError as e:
    check("坏文件(解析失败): 抛 ValueError(入库失败: ...)", "入库失败" in str(e), str(e))

# ---------- 8) clear_collection：清空该师 collection ----------
store.clear_collection(TID)
check("清库: collection_count==0", store.collection_count(TID) == 0)
check("清库: list_documents 为空", store.list_documents(TID) == [])
check("清库: 缓存已失效", all(not str(k).startswith(f"retr|{TID}|") for k in store.retrieve_cache._data))

print("\nRESULT:", "ALL PASS" if ok else "HAS FAILURES")
sys.exit(0 if ok else 1)
