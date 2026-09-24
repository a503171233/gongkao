# -*- coding: utf-8 -*-
"""课件入库：提取文本 → 清洗 → 分块。对应说明书 §6.1/§6.2。"""
import re
from dataclasses import dataclass, field
from pathlib import Path

from .config import Config


@dataclass
class Chunk:
    content: str
    doc_name: str
    idx: int
    tags: str = ""


def extract_text(path: str | Path) -> str:
    """按扩展名提取课件文本。解析失败抛 ValueError("入库失败: ...")，不静默入库。"""
    p = Path(path)
    ext = p.suffix.lower()
    if ext in (".md", ".txt", ".text"):
        try:
            return p.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as e:
            raise ValueError(f"入库失败: 文件不是有效 UTF-8 编码: {p.name}") from e
    if ext == ".docx":
        import docx  # python-docx

        d = docx.Document(str(p))
        return "\n".join(par.text for par in d.paragraphs if par.text.strip())
    if ext == ".pptx":
        from pptx import Presentation  # python-pptx

        prs = Presentation(str(p))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t:
                        parts.append(t)
        return "\n".join(parts)
    if ext == ".pdf":
        import fitz  # pymupdf

        doc = fitz.open(str(p))
        return "\n".join(page.get_text("text") for page in doc)
    if ext == ".xlsx":
        import openpyxl  # #24 R3 知识体系文档建树支持 Excel

        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        parts = []
        try:
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        finally:
            wb.close()
        return "\n".join(parts)
    if ext == ".xls":
        raise ValueError("入库失败: 暂不支持旧版 .xls（请用 Excel 另存为 .xlsx 后重传）")
    raise ValueError(f"入库失败: 暂不支持该格式: {ext}（支持 md/txt/docx/pptx/pdf/xlsx）")


# ---------- #33 R2：AI 采集专用解析（文本 + 图片内联标记） ----------

# 合法图片 URL（/api/qimg/{teacher_id}/{sha1前12}.{ext}）——qimg 端点白名单与此一致
QIMG_URL_RE = re.compile(
    r"^/api/qimg/[A-Za-z0-9_-]+/[a-f0-9]{12}\.(?:png|jpe?g|gif|webp)$")
# 文本中的 Markdown 图片标记：![图N](/api/qimg/...)
IMG_MARK_RE = re.compile(r"!\[(图\d+)\]\((/api/qimg/[^)\s]+)\)")


def _img_ext(blob: bytes) -> str:
    """按文件头嗅探图片扩展名（内容寻址落盘用）。"""
    if blob[:8] == b"\x89PNG\r\n\x1a\n":
        return ".png"
    if blob[:3] == b"\xff\xd8\xff":
        return ".jpg"
    if blob[:6] in (b"GIF87a", b"GIF89a"):
        return ".gif"
    if blob[:4] == b"RIFF" and blob[8:12] == b"WEBP":
        return ".webp"
    return ".png"


def _save_img(img_dir: Path, blob: bytes) -> str:
    """图片内容寻址落盘：sha1 前 12 位命名，同图重复出现天然幂等。返回文件名。"""
    import hashlib

    fname = hashlib.sha1(blob).hexdigest()[:12] + _img_ext(blob)
    fp = img_dir / fname
    if not fp.exists():
        img_dir.mkdir(parents=True, exist_ok=True)
        fp.write_bytes(blob)
    return fname


def _img_mark(n: int, teacher_id: str, fname: str) -> str:
    return f"![图{n}](/api/qimg/{teacher_id}/{fname})"


def _docx_collect(p: Path, img_dir: Path, tid: str) -> tuple[str, int]:
    """docx：段落文本 + 段落内 drawing 图片（w:drawing → a:blip r:embed → rId → blob）。"""
    import docx
    from docx.oxml.ns import qn

    d = docx.Document(str(p))
    parts: list[str] = []
    n = 0
    for par in d.paragraphs:
        t = par.text.strip()
        if t:
            parts.append(t)
        for blip in par._element.findall(".//" + qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            if not rid:
                continue
            part = d.part.related_parts.get(rid)
            try:
                blob = part.blob if part is not None else b""
            except Exception:  # noqa: BLE001 单图损坏不阻断整篇采集
                blob = b""
            if not blob:
                continue
            n += 1
            parts.append(_img_mark(n, tid, _save_img(img_dir, blob)))
    return "\n".join(parts), n


def _pptx_collect(p: Path, img_dir: Path, tid: str) -> tuple[str, int]:
    """pptx：文本 shape + 图片 shape（shape_type==13 PICTURE → shape.image.blob）。"""
    from pptx import Presentation

    prs = Presentation(str(p))
    parts: list[str] = []
    n = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                try:
                    blob = shape.image.blob
                except Exception:  # noqa: BLE001 单图损坏不阻断
                    continue
                if not blob:
                    continue
                n += 1
                parts.append(_img_mark(n, tid, _save_img(img_dir, blob)))
            elif shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    parts.append(t)
    return "\n".join(parts), n


def _pdf_collect(p: Path, img_dir: Path, tid: str) -> tuple[str, int]:
    """pdf：get_text('dict') 按文档流顺序混合输出文本块与图片块（type==1）。"""
    import fitz

    doc = fitz.open(str(p))
    parts: list[str] = []
    n = 0
    for page in doc:
        blocks = page.get_text("dict").get("blocks", [])
        for blk in blocks:
            if blk.get("type") == 1:  # image block
                blob = blk.get("image") or b""
                if not blob:
                    continue
                n += 1
                parts.append(_img_mark(n, tid, _save_img(img_dir, blob)))
            else:  # text block：行拼文本（保持块内行结构）
                lines = []
                for ln in blk.get("lines", []):
                    s = "".join(sp.get("text", "") for sp in ln.get("spans", []))
                    if s.strip():
                        lines.append(s.strip())
                if lines:
                    parts.append("\n".join(lines))
    return "\n".join(parts), n


def extract_img_urls(text: str) -> list[str]:
    """从文本提取合法图片 URL 列表（去重保序；入库 images 列用）。#33 R2"""
    out: list[str] = []
    for m in re.finditer(r"!\[[^\]]*\]\((/api/qimg/[^)\s]+)\)", text):
        u = m.group(1)
        if QIMG_URL_RE.match(u) and u not in out:
            out.append(u)
    return out


def extract_collect(path: str | Path, teacher_id: str = "") -> dict:
    """AI 采集专用解析：提取文本并按文档流顺序内联图片标记（#33 R2）。

    - 不改 extract_text 签名/行为，文档入库与向量链路零影响；
    - 图片内容寻址落盘 data/qimg/{teacher_id}/（sha1 前 12 位，同图幂等）；
    - 文本中插入 Markdown 标记 ![图N](/api/qimg/{tid}/{fname})，
      后续 LLM 提取时原样保留进题干，前端按白名单渲染；
    - 返回 {text, images: [{marker, url}], n_img}；
    - 解析失败抛 ValueError（消息以"解析失败:"开头）。
    """
    from .config import Config

    cfg = Config()
    p = Path(path)
    ext = p.suffix.lower()
    tid = (teacher_id or "").strip() or "common"
    img_dir = cfg.data_dir / "qimg" / tid
    text = ""
    n = 0
    if ext in (".md", ".txt", ".text"):
        try:
            text = p.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as e:
            raise ValueError(f"解析失败: 文件不是有效 UTF-8 编码: {p.name}") from e
    elif ext == ".docx":
        text, n = _docx_collect(p, img_dir, tid)
    elif ext == ".pptx":
        text, n = _pptx_collect(p, img_dir, tid)
    elif ext == ".pdf":
        text, n = _pdf_collect(p, img_dir, tid)
    else:
        raise ValueError(
            f"解析失败: 暂不支持该格式: {ext or '(无扩展名)'}（支持 md/txt/docx/pptx/pdf）")
    images = [{"marker": m.group(1), "url": m.group(2)}
              for m in IMG_MARK_RE.finditer(text)]
    return {"text": text, "images": images, "n_img": len(images)}


def clean_text(text: str) -> str:
    """轻量清洗：去空行/多余空白，保留纯干货（说明书 §6.1 清洗）。"""
    text = re.sub(r"\r\n?", "\n", text)
    lines = [ln.strip() for ln in text.split("\n")]
    lines = [ln for ln in lines if ln]
    return "\n".join(lines)


def chunk_text(text: str, chunk_size: int = 600, overlap: int = 120) -> list[str]:
    """按段落自然边界聚合分块，字符数兜底；块间重叠 overlap 字符（§6.2）。"""
    text = clean_text(text)
    paras = text.split("\n")
    chunks: list[str] = []
    buf = ""
    for para in paras:
        if len(buf) + len(para) + 1 <= chunk_size:
            buf = f"{buf}\n{para}".strip()
            continue
        # 当前缓冲已满：先提交
        if buf:
            chunks.append(buf)
            tail = buf[-overlap:] if len(buf) >= overlap else buf
            buf = f"{tail}\n{para}".strip()
        else:
            buf = para
        # 单段超长时硬切
        while len(buf) > chunk_size:
            head, buf = buf[:chunk_size], buf[chunk_size:]
            chunks.append(head)
    if buf:
        chunks.append(buf)
    return chunks


def ingest_file(path: str | Path, cfg: Config) -> list[Chunk]:
    """解析单个课件文件为分块列表（带来源与序号元数据）。
    空文件 / 不支持格式 / 解析失败一律抛 ValueError("入库失败: ...")，不静默入库。
    """
    p = Path(path)
    try:
        text = extract_text(p)
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"入库失败: 解析文件失败({p.name}): {e}") from e
    if not text.strip():
        raise ValueError(f"入库失败: 文件未提取到文本: {p.name}")
    parts = chunk_text(text, cfg.chunk_size, cfg.chunk_overlap)
    return [
        Chunk(content=part, doc_name=p.name, idx=i) for i, part in enumerate(parts)
    ]
